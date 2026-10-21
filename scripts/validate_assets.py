#!/usr/bin/env python3
"""
Enterprise Power BI Analytics Portfolio — Automated Asset Validation Engine
Validates dashboard images, SVG icon syntax, directory structure, and Excel handbook.
"""

import os
import sys
import xml.etree.ElementTree as ET
from PIL import Image
import openpyxl

def run_asset_validation():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    repo_root = os.path.abspath(os.path.join(script_dir, '..'))

    print("================================================================================")
    print("ENTERPRISE POWER BI ANALYTICS PORTFOLIO — ASSET VALIDATION")
    print(f"Repository Root: {repo_root}")
    print("================================================================================")

    failures = []

    # 1. Dashboard Images Validation
    dash_dir = os.path.join(repo_root, 'assets', 'dashboards')
    expected_pages = [
        '01-executive-overview',
        '02-sales-performance',
        '03-product-analysis',
        '04-customer-insights',
        '05-region-analysis',
        '06-order-details',
        '07-profitability-analysis',
        '08-return-analysis',
        '09-forecast-analysis'
    ]

    print("\n--------------------------------------------------------------------------------")
    print("DASHBOARD IMAGE ASSETS (PNG & WEBP)")
    print("--------------------------------------------------------------------------------")

    for stem in expected_pages:
        for ext in ['.png', '.webp']:
            fname = f"{stem}{ext}"
            fpath = os.path.join(dash_dir, fname)
            if not os.path.exists(fpath):
                failures.append(f"Missing dashboard image: {fname}")
                print(f"[FAIL] Missing {fname}")
                continue
            try:
                with Image.open(fpath) as img:
                    w, h = img.size
                    if w > 0 and h > 0:
                        print(f"[PASS] {fname}: {img.format} {w}x{h} ({os.path.getsize(fpath):,} bytes)")
                    else:
                        failures.append(f"Invalid image dimensions for {fname}: {w}x{h}")
                        print(f"[FAIL] {fname}: Invalid dimensions")
            except Exception as e:
                failures.append(f"Corrupt image {fname}: {str(e)}")
                print(f"[FAIL] {fname}: {e}")

    # 2. SVG Icon Library Validation
    icons_dir = os.path.join(repo_root, 'assets', 'icons')
    print("\n--------------------------------------------------------------------------------")
    print("SVG ICON VECTOR ASSETS (XML SYNTAX & SECURITY)")
    print("--------------------------------------------------------------------------------")

    svg_count = 0
    for root, dirs, files in os.walk(icons_dir):
        for f in files:
            if f.endswith('.svg'):
                svg_count += 1
                fpath = os.path.join(root, f)
                rel_path = os.path.relpath(fpath, icons_dir)
                try:
                    # XML parse check
                    tree = ET.parse(fpath)
                    root_elem = tree.getroot()
                    if not root_elem.tag.endswith('svg'):
                        failures.append(f"File {rel_path} root tag is not svg ({root_elem.tag})")
                        print(f"[FAIL] {rel_path}: root tag is not svg")
                        continue

                    # Security check: ensure no embedded <script> tags
                    with open(fpath, 'r', encoding='utf-8') as sf:
                        content = sf.read().lower()
                        if '<script' in content or 'javascript:' in content or 'onload=' in content:
                            failures.append(f"Security risk detected in {rel_path}: script or event handler found")
                            print(f"[FAIL] {rel_path}: script or event handler detected")
                            continue

                except ET.ParseError as pe:
                    failures.append(f"Invalid XML syntax in {rel_path}: {str(pe)}")
                    print(f"[FAIL] {rel_path}: XML parse error: {pe}")
                except Exception as ex:
                    failures.append(f"Error checking {rel_path}: {str(ex)}")
                    print(f"[FAIL] {rel_path}: {ex}")

    print(f"[INFO] Audited {svg_count} SVG vector icon files.")
    if svg_count >= 60:
        print(f"[PASS] SVG Icon count ({svg_count}) satisfies requirement (>= 60)")
    else:
        failures.append(f"Insufficient SVG icons: found {svg_count}, expected >= 60")
        print(f"[FAIL] SVG Icon count: {svg_count}")

    # 3. Excel Handbook Validation
    excel_path = os.path.join(repo_root, 'excel', 'powerbi-implementation-handbook.xlsx')
    print("\n--------------------------------------------------------------------------------")
    print("EXCEL IMPLEMENTATION HANDBOOK AUDIT")
    print("--------------------------------------------------------------------------------")
    if os.path.exists(excel_path):
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True)
            print(f"[PASS] Excel handbook readable: {len(wb.sheetnames)} sheets found")
            expected_sheets = ['START HERE', 'DATA DICTIONARY', 'MODEL RELATIONSHIPS', 'DAX MEASURE LIBRARY', 'VISUAL BUILD MATRIX']
            for s in expected_sheets:
                if s in wb.sheetnames:
                    print(f"  [PASS] Sheet '{s}' verified")
                else:
                    failures.append(f"Missing expected sheet '{s}' in handbook")
                    print(f"  [FAIL] Missing sheet '{s}'")
        except Exception as e:
            failures.append(f"Error opening Excel handbook: {str(e)}")
            print(f"[FAIL] Error opening Excel handbook: {e}")
    else:
        failures.append("Missing Excel handbook at excel/powerbi-implementation-handbook.xlsx")
        print("[FAIL] Missing excel/powerbi-implementation-handbook.xlsx")

    # 4. Branding Logo Validation
    branding_logo = os.path.join(repo_root, 'assets', 'branding', 'datta-sable-monogram.svg')
    if os.path.exists(branding_logo):
        print(f"[PASS] Branding monogram exists: {branding_logo}")
    else:
        failures.append("Missing branding monogram at assets/branding/datta-sable-monogram.svg")
        print("[FAIL] Missing branding monogram")

    print("\n================================================================================")
    if not failures:
        print("ALL ASSET VALIDATION CHECKS PASSED")
        print("================================================================================")
        return 0
    else:
        print(f"FAILED ASSET CHECKS: {len(failures)}")
        for f in failures:
            print(f"  - {f}")
        print("================================================================================")
        return 1

if __name__ == '__main__':
    sys.exit(run_asset_validation())
